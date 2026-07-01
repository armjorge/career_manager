from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any

from docx import Document


def _format_value(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return None
    return text.replace("\\n", "\n")


def format_date_issued(lang: str | None, app_date: date | datetime | None) -> str:
    if app_date is None:
        return ""

    if isinstance(app_date, datetime):
        dt = app_date
    else:
        dt = datetime.combine(app_date, datetime.min.time())

    day = dt.day
    month_num = dt.month
    year = dt.year

    months_en = [
        "January",
        "February",
        "March",
        "April",
        "May",
        "June",
        "July",
        "August",
        "September",
        "October",
        "November",
        "December",
    ]
    months_es = [
        "enero",
        "febrero",
        "marzo",
        "abril",
        "mayo",
        "junio",
        "julio",
        "agosto",
        "septiembre",
        "octubre",
        "noviembre",
        "diciembre",
    ]
    months_fr = [
        "janvier",
        "février",
        "mars",
        "avril",
        "mai",
        "juin",
        "juillet",
        "août",
        "septembre",
        "octobre",
        "novembre",
        "décembre",
    ]

    if lang == "English":
        suffix = "th"
        if day in (1, 21, 31):
            suffix = "st"
        elif day in (2, 22):
            suffix = "nd"
        elif day in (3, 23):
            suffix = "rd"
        return f"Mexico City, {months_en[month_num - 1]} {day}{suffix}, {year}"

    if lang == "Spanish":
        return f"Ciudad de México, {day} de {months_es[month_num - 1].capitalize()} de {year}"

    if lang == "French":
        return f"Mexico, le {day} {months_fr[month_num - 1].capitalize()} {year}"

    return ""


def _replace_placeholder_in_paragraph(
    doc: Document,
    paragraph,
    placeholder: str,
    replacement: str | None,
    *,
    multiline: bool,
) -> None:
    if placeholder not in paragraph.text:
        return

    value = replacement or ""
    if not multiline or "\n" not in value:
        paragraph.text = paragraph.text.replace(placeholder, value)
        return

    parts = value.split("\n")
    paragraph.text = paragraph.text.replace(placeholder, parts[0])

    p_element = paragraph._element
    body_element = doc._body._element
    index = list(body_element).index(p_element)

    for part in parts[1:]:
        new_p = doc.add_paragraph(part, style=paragraph.style.name)
        new_p_element = new_p._element
        body_element.remove(new_p_element)
        body_element.insert(index + 1, new_p_element)
        index += 1


def populate_document(template_bytes: bytes, row: dict[str, Any]) -> bytes:
    doc = Document(io.BytesIO(template_bytes))
    placeholders = {f"{{{key}}}": _format_value(value) for key, value in row.items()}

    for paragraph in doc.paragraphs:
        for placeholder, value in placeholders.items():
            _replace_placeholder_in_paragraph(doc, paragraph, placeholder, value, multiline=True)

    for table in doc.tables:
        for table_row in table.rows:
            for cell in table_row.cells:
                for paragraph in cell.paragraphs:
                    for placeholder, value in placeholders.items():
                        _replace_placeholder_in_paragraph(doc, paragraph, placeholder, value, multiline=False)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def propose_output_filename(
    prefix: str | None,
    company_name: str,
    category_name: str | None,
    category: str,
) -> str:
    cat_suffix = "CV" if category == "Resume" else "CLetter"
    role = category_name if category_name else cat_suffix
    parts = [part.strip() for part in (prefix or "", company_name or "", role) if part.strip()]
    return f"{' '.join(parts)}.docx"


def unique_output_filename(base_name: str, existing_names: set[str]) -> str:
    if base_name not in existing_names:
        return base_name

    if not base_name.lower().endswith(".docx"):
        stem = base_name
        ext = ""
    else:
        stem = base_name[:-5]
        ext = ".docx"

    counter = 1
    while True:
        candidate = f"{stem}_{counter}{ext}"
        if candidate not in existing_names:
            return candidate
        counter += 1
